"""Load stage: SQLite storage (idempotent) + formatted Excel report."""

import logging
import sqlite3

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from etl_pipeline import transform

log = logging.getLogger("etl.load")

# --- Excel styling ---------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
TOTAL_FONT = Font(bold=True, size=11)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

REPORT_COLUMNS = [
    ("capital", "Capital", 14, None),
    ("country_name", "Country", 16, None),
    ("region", "Region", 22, None),
    ("income_level", "Income level", 14, None),
    ("population_2024", "Population (2024)", 17, "#,##0"),
    ("local_time", "Local time", 17, None),
    ("temp_c", "Temp °C", 10, "0.0"),
    ("weather_label", "Weather", 20, None),
    ("humidity_pct", "Humidity %", 12, "0"),
    ("wind_kmh", "Wind km/h", 12, "0.0"),
    ("currency", "Currency", 10, None),
    ("usd_to_currency", "USD → CUR", 12, "0.0000"),
    ("fx_date", "FX date", 12, None),
]


def init_db(db_path=None, schema_path=None):
    """Create tables/view from schema.sql (safe to run repeatedly)."""
    db_path = db_path or config.DB_PATH
    schema_path = schema_path or config.SCHEMA_PATH
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    with open(schema_path, encoding="utf-8") as fh:
        conn.executescript(fh.read())
    conn.commit()
    return conn


def upsert_countries(conn, rows):
    conn.executemany(
        """INSERT INTO countries
           (iso3, country_name, region, income_level, capital,
            population_2024, latitude, longitude, timezone, updated_at)
           VALUES (:iso3, :country_name, :region, :income_level, :capital,
                   :population_2024, :latitude, :longitude, :timezone, :updated_at)
           ON CONFLICT(iso3) DO UPDATE SET
             country_name=excluded.country_name, region=excluded.region,
             income_level=excluded.income_level, capital=excluded.capital,
             population_2024=excluded.population_2024,
             latitude=excluded.latitude, longitude=excluded.longitude,
             timezone=excluded.timezone, updated_at=excluded.updated_at""",
        rows,
    )
    conn.commit()
    log.info("countries upserted: %d", len(rows))


def replace_weather(conn, run_id, rows):
    conn.executemany(
        """INSERT OR REPLACE INTO weather_observations
           (run_id, iso3, observed_at_utc, temp_c, humidity_pct, wind_kmh,
            weather_code, weather_label, utc_offset_seconds)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        [(run_id, r["iso3"], r["observed_at_utc"], r["temp_c"],
          r["humidity_pct"], r["wind_kmh"], r["weather_code"],
          r["weather_label"], r["utc_offset_seconds"]) for r in rows],
    )
    conn.commit()
    log.info("weather rows replaced for run %s: %d", run_id, len(rows))


def replace_fx(conn, run_id, rows):
    conn.executemany(
        """INSERT OR REPLACE INTO fx_rates
           (run_id, iso3, currency, usd_to_currency, fx_date)
           VALUES (?, ?, ?, ?, ?)""",
        [(run_id, r["iso3"], r["currency"], r["usd_to_currency"], r["fx_date"])
         for r in rows],
    )
    conn.commit()
    log.info("fx rows replaced for run %s: %d", run_id, len(rows))


def _style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=2, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def export_excel(db_path, xlsx_path, run_id):
    """Build the formatted workbook from the reporting view."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT * FROM v_capital_report WHERE run_id = ? ORDER BY capital",
        (run_id,),
    ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Capitals"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.merge_cells("A1:M1")
    title = ws["A1"]
    title.value = f"World Capitals Dashboard — live data snapshot {run_id}"
    title.font = TITLE_FONT

    headers = [label for _, label, _, _ in REPORT_COLUMNS]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=2, column=col, value=header)
    _style_header(ws, len(headers))

    data_row_count = 0
    for i, r in enumerate(rows, start=3):
        d = dict(r)
        d["local_time"] = transform.local_time_str(
            d.get("observed_at_utc"), d.get("utc_offset_seconds"))
        for col, (key, _label, _w, numfmt) in enumerate(REPORT_COLUMNS, start=1):
            cell = ws.cell(row=i, column=col, value=d.get(key))
            cell.border = BORDER
            if numfmt and isinstance(cell.value, (int, float)):
                cell.number_format = numfmt
        data_row_count += 1

    # Totals / averages row
    total_row = 3 + data_row_count
    ws.cell(row=total_row, column=1, value="TOTAL / AVG").font = TOTAL_FONT
    pop_col = next(i for i, (k, *_r) in enumerate(REPORT_COLUMNS, 1)
                   if k == "population_2024")
    temp_col = next(i for i, (k, *_r) in enumerate(REPORT_COLUMNS, 1)
                    if k == "temp_c")
    hum_col = next(i for i, (k, *_r) in enumerate(REPORT_COLUMNS, 1)
                   if k == "humidity_pct")
    wind_col = next(i for i, (k, *_r) in enumerate(REPORT_COLUMNS, 1)
                    if k == "wind_kmh")
    if data_row_count:
        first, last = 3, total_row - 1
        specs = [(pop_col, "SUM", "#,##0"), (temp_col, "AVERAGE", "0.0"),
                 (hum_col, "AVERAGE", "0"), (wind_col, "AVERAGE", "0.0")]
        for col, func, numfmt in specs:
            letter = get_column_letter(col)
            cell = ws.cell(row=total_row, column=col,
                           value=f"={func}({letter}{first}:{letter}{last})")
            cell.number_format = numfmt
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = BORDER
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).fill = TOTAL_FILL
        ws.cell(row=total_row, column=col).border = BORDER

    for col, (_k, _l, width, _f) in enumerate(REPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{total_row}"
    ws.freeze_panes = "A3"

    # Bar chart: temperature per capital
    if data_row_count:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Current temperature by capital (°C)"
        chart.y_axis.title = "°C"
        data = Reference(ws, min_col=temp_col, min_row=2,
                         max_row=total_row - 1)
        cats = Reference(ws, min_col=1, min_row=3, max_row=total_row - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 9, 18
        ws.add_chart(chart, "O3")

    # --- Region summary sheet + pie chart ----------------------------------
    ws2 = wb.create_sheet("By Region")
    regions = conn.execute(
        """SELECT c.region,
                  COUNT(*) AS capitals,
                  SUM(c.population_2024) AS population,
                  AVG(w.temp_c) AS avg_temp
           FROM countries c
           LEFT JOIN weather_observations w
             ON w.iso3 = c.iso3 AND w.run_id = ?
           GROUP BY c.region ORDER BY population DESC""",
        (run_id,),
    ).fetchall()
    ws2.merge_cells("A1:D1")
    ws2["A1"] = f"Summary by region — {run_id}"
    ws2["A1"].font = TITLE_FONT
    for col, header in enumerate(
            ["Region", "Capitals", "Population (2024)", "Avg temp °C"], 1):
        cell = ws2.cell(row=2, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    for i, (region, capitals, population, avg_temp) in enumerate(regions, 3):
        for col, value in enumerate(
                [region, capitals, population,
                 round(avg_temp, 1) if avg_temp is not None else None], 1):
            cell = ws2.cell(row=i, column=col, value=value)
            cell.border = BORDER
        ws2.cell(row=i, column=3).number_format = "#,##0"
    for col, width in zip("ABCD", (28, 10, 20, 14)):
        ws2.column_dimensions[col].width = width
    if regions:
        pie = PieChart()
        pie.title = "Population share by region"
        labels = Reference(ws2, min_col=1, min_row=3, max_row=2 + len(regions))
        data = Reference(ws2, min_col=3, min_row=2, max_row=2 + len(regions))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height, pie.width = 9, 16
        ws2.add_chart(pie, "F3")

    conn.close()
    wb.save(xlsx_path)
    log.info("Excel report written: %s (%d capital rows)", xlsx_path,
             data_row_count)
    return data_row_count
